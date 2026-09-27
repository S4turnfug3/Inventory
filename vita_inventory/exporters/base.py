"""
Basis-Exporter für verschiedene Formate.
"""

from abc import ABC, abstractmethod
from pathlib import Path
from typing import Dict, Any, List
from vita_inventory.core.logger import setup_logger

logger = setup_logger(__name__)


class BaseExporter(ABC):
    """Basis-Klasse für Exporter."""
    
    def __init__(self, output_dir: str = 'output'):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
    
    @abstractmethod
    def export(self, data: Dict[str, Any], filename: str) -> Path:
        """Exportiert Daten."""
        pass


class JSONExporter(BaseExporter):
    """JSON-Exporter."""
    
    def export(self, data: Dict[str, Any], filename: str = 'inventory.json') -> Path:
        """Exportiert zu JSON."""
        import json
        
        output_path = self.output_dir / filename
        try:
            with open(output_path, 'w', encoding='utf-8') as f:
                json.dump(data, f, indent=2, ensure_ascii=False)
            logger.info(f"JSON exportiert: {output_path}")
            return output_path
        except Exception as e:
            logger.error(f"JSON-Export fehlgeschlagen: {e}")
            raise


class MarkdownExporter(BaseExporter):
    """Markdown-Exporter."""
    
    def export(self, data: Dict[str, Any], filename: str = 'inventory.md') -> Path:
        """Exportiert zu Markdown."""
        output_path = self.output_dir / filename
        
        try:
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write("# Vita Inventory Report\n\n")
                
                # System-Infos
                if data.get('system'):
                    f.write("## System Information\n\n")
                    system = data['system']
                    f.write(f"- **Hostname**: {system.get('hostname', 'N/A')}\n")
                    f.write(f"- **OS**: {system.get('os_name', 'N/A')} {system.get('os_version', '')}\n")
                    f.write(f"- **Architecture**: {system.get('architecture', 'N/A')}\n")
                    f.write(f"- **CPU**: {system.get('cpu_count', 'N/A')} cores\n")
                    if system.get('cpu_model'):
                        f.write(f"- **CPU Model**: {system['cpu_model']}\n")
                    f.write(f"- **Memory**: {system.get('memory_gb', 0):.2f} GB\n")
                    if system.get('manufacturer'):
                        f.write(f"- **Manufacturer**: {system['manufacturer']}\n")
                    if system.get('serial_number'):
                        f.write(f"- **Serial**: {system['serial_number']}\n")
                    f.write("\n")
                
                # Netzwerk-Infos
                if data.get('network'):
                    f.write("## Network Information\n\n")
                    network = data['network']
                    f.write(f"- **Interface**: {network.get('interface', 'N/A')}\n")
                    if network.get('ipv4_address'):
                        f.write(f"- **IPv4**: {network['ipv4_address']}/{network.get('ipv4_prefix', '24')}\n")
                        f.write(f"- **Network**: {network.get('ipv4_network', 'N/A')}\n")
                    if network.get('ipv4_gateway'):
                        f.write(f"- **Gateway**: {network['ipv4_gateway']}\n")
                    if network.get('ipv6_address'):
                        f.write(f"- **IPv6**: {network['ipv6_address']}/{network.get('ipv6_prefix', '64')}\n")
                    if network.get('mac_address'):
                        f.write(f"- **MAC**: {network['mac_address']}\n")
                    f.write("\n")
                
                # Geräte
                if data.get('devices'):
                    f.write("## Network Devices\n\n")
                    f.write("| IP Address | MAC Address | Status | Hostname |\n")
                    f.write("|---|---|---|---|\n")
                    for device in data['devices']:
                        ip = device.get('ipv4_address', 'N/A')
                        mac = device.get('mac_address', 'N/A')
                        status = device.get('status', 'unknown')
                        hostname = device.get('hostname') or 'N/A'
                        f.write(f"| {ip} | {mac} | {status} | {hostname} |\n")
                    f.write(f"\n**Total Devices**: {len(data['devices'])}\n")
            
            logger.info(f"Markdown exportiert: {output_path}")
            return output_path
        
        except Exception as e:
            logger.error(f"Markdown-Export fehlgeschlagen: {e}")
            raise


class CSVExporter(BaseExporter):
    """CSV-Exporter."""
    
    def export(self, data: Dict[str, Any], filename: str = 'inventory.csv') -> Path:
        """Exportiert zu CSV."""
        import csv
        
        output_path = self.output_dir / filename
        
        try:
            with open(output_path, 'w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f)
                
                # Header
                writer.writerow([
                    'IP Address', 'MAC Address', 'Hostname', 'Status',
                    'System Hostname', 'OS', 'CPU Cores', 'Memory GB'
                ])
                
                devices = data.get('devices', [])
                system = data.get('system', {})
                
                # Zeilen
                for device in devices:
                    writer.writerow([
                        device.get('ipv4_address', ''),
                        device.get('mac_address', ''),
                        device.get('hostname', ''),
                        device.get('status', ''),
                        system.get('hostname', ''),
                        f"{system.get('os_name', '')} {system.get('os_version', '')}",
                        system.get('cpu_count', ''),
                        f"{system.get('memory_gb', 0):.2f}"
                    ])
            
            logger.info(f"CSV exportiert: {output_path}")
            return output_path
        
        except Exception as e:
            logger.error(f"CSV-Export fehlgeschlagen: {e}")
            raise


class ExcelExporter(BaseExporter):
    """Excel-Exporter (optional, mit openpyxl)."""
    
    def export(self, data: Dict[str, Any], filename: str = 'inventory.xlsx') -> Path:
        """Exportiert zu Excel."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            logger.warning("openpyxl nicht installiert. Nutze stattdessen CSV.")
            csv_exporter = CSVExporter(str(self.output_dir))
            return csv_exporter.export(data, filename.replace('.xlsx', '.csv'))
        
        output_path = self.output_dir / filename
        
        try:
            wb = openpyxl.Workbook()
            
            # System-Sheet
            if data.get('system'):
                ws_system = wb.active
                ws_system.title = "System"
                system = data['system']
                
                header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                header_font = Font(bold=True, color="FFFFFF")
                
                headers = ['Parameter', 'Value']
                ws_system.append(headers)
                for cell in ws_system[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                
                ws_system.append(['Hostname', system.get('hostname', '')])
                ws_system.append(['OS', f"{system.get('os_name', '')} {system.get('os_version', '')}"])
                ws_system.append(['Architecture', system.get('architecture', '')])
                ws_system.append(['CPU Cores', system.get('cpu_count', '')])
                ws_system.append(['CPU Model', system.get('cpu_model', '')])
                ws_system.append(['Memory (GB)', f"{system.get('memory_gb', 0):.2f}"])
                ws_system.append(['Manufacturer', system.get('manufacturer', '')])
                ws_system.append(['Serial Number', system.get('serial_number', '')])
                
                ws_system.column_dimensions['A'].width = 20
                ws_system.column_dimensions['B'].width = 40
            
            # Netzwerk-Sheet
            if data.get('network'):
                ws_network = wb.create_sheet("Network")
                network = data['network']
                
                ws_network.append(['Parameter', 'Value']
                for cell in ws_network[1]:
                    cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
                    cell.font = Font(bold=True, color="FFFFFF")
                
                ws_network.append(['Interface', network.get('interface', '')])
                ws_network.append(['IPv4 Address', network.get('ipv4_address', '')])
                ws_network.append(['IPv4 Prefix', network.get('ipv4_prefix', '')])
                ws_network.append(['IPv4 Network', network.get('ipv4_network', '')])
                ws_network.append(['Gateway', network.get('ipv4_gateway', '')])
                ws_network.append(['MAC Address', network.get('mac_address', '')])
                if network.get('ipv6_address'):
                    ws_network.append(['IPv6 Address', network['ipv6_address']])
                    ws_network.append(['IPv6 Prefix', network.get('ipv6_prefix', '')])
                
                ws_network.column_dimensions['A'].width = 20
                ws_network.column_dimensions['B'].width = 40
            
            # Geräte-Sheet
            if data.get('devices'):
                ws_devices = wb.create_sheet("Devices")
                
                headers = ['IP Address', 'MAC Address', 'Hostname', 'Status']
                ws_devices.append(headers)
                for cell in ws_devices[1]:
                    cell.fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
                    cell.font = Font(bold=True)
                
                for device in data['devices']:
                    ws_devices.append([
                        device.get('ipv4_address', ''),
                        device.get('mac_address', ''),
                        device.get('hostname', ''),
                        device.get('status', '')
                    ])
                
                ws_devices.column_dimensions['A'].width = 15
                ws_devices.column_dimensions['B'].width = 20
                ws_devices.column_dimensions['C'].width = 25
                ws_devices.column_dimensions['D'].width = 12
            
            wb.save(output_path)
            logger.info(f"Excel exportiert: {output_path}")
            return output_path
        
        except Exception as e:
            logger.error(f"Excel-Export fehlgeschlagen: {e}")
            raise
